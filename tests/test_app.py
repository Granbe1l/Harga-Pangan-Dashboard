"""Navigation, calendar comparisons, empty states, and report exports."""
from pathlib import Path
from io import BytesIO
import unittest
import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

ROOT=Path(__file__).resolve().parents[1]
TITLE='Analisis Tren dan Perbandingan Harga Pangan Antarprovinsi'
PAGES=['Ringkasan','Tren & fluktuasi','Perbandingan provinsi']

class DashboardTests(unittest.TestCase):
    def open_app(self):
        app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=45).run()
        self.good(app)
        return app

    def good(self,app):
        self.assertEqual([e.message for e in app.exception],[])

    def menu(self,app,name):
        app.radio(key='page').set_value(name).run();self.good(app)

    def test_three_analysis_pages_render_and_have_sources(self):
        app=self.open_app()
        self.assertEqual(app.radio(key='page').options,PAGES)
        for page in PAGES:
            self.menu(app,page)
            self.assertEqual(len(app.metric),3)
            self.assertTrue(any(e.label=='Sumber & cara menghitung' for e in app.expander))
            text=' '.join(m.value for m in app.markdown)
            self.assertNotIn('Keranjang',text)
            self.assertFalse(app.get('data_editor'))

    def test_default_findings_match_observed_prices(self):
        app=self.open_app()
        self.assertEqual(app.title[0].value,TITLE)
        self.assertEqual([m.value for m in app.metric],['25 / 25','2','+11,16%'])
        self.assertIn('Cabai Merah Keriting',app.metric[2].label)
        self.assertTrue(any('Daging Kerbau Segar (Lokal)' in c.value and 'tidak tersedia' in c.value for c in app.caption))
        self.assertFalse(any(s.key=='filter_threshold' for s in app.slider))
        self.assertTrue(any('4 dari 12 pasangan bulan' in m.value for m in app.markdown))
        self.assertTrue(any('Studi historis Januari 2021–Oktober 2025' in c.value for c in app.caption))

    def test_first_month_has_unknown_changes(self):
        app=self.open_app()
        app.selectbox(key='filter_as_of').set_value(pd.Timestamp('2021-01-01')).run();self.good(app)
        self.assertEqual(app.metric[1].value,'—')
        self.assertEqual(app.metric[2].value,'—')
        self.assertEqual(len(app.get('plotly_chart')),0)

    def kpi_details(self, app):
        availability = next(e for e in app.expander if e.label == 'Lihat rincian ketersediaan harga komoditas')
        rises = next(e for e in app.expander if e.label == 'Lihat komoditas dengan kenaikan ≥10%')
        self.assertFalse(availability.proto.expanded)
        self.assertFalse(rises.proto.expanded)
        return availability, rises

    def test_kpi_details_follow_province_month_and_window(self):
        from analyst.pipeline import load_clean_data
        from analyst.analytics import comparison_frame, monthly_findings
        changes = comparison_frame(load_clean_data(ROOT))
        app = self.open_app()
        for province, month, available, count in [
            ('Jawa Barat', '2025-10-01', 25, 2),
            ('DI Yogyakarta', '2025-10-01', 23, 1),
            ('Jawa Barat', '2021-02-01', 12, 0),
            ('Jawa Barat', '2021-01-01', 12, None),
            ('Papua Pegunungan', '2025-10-01', 0, None),
        ]:
            with self.subTest(province=province, month=month):
                app.selectbox(key='filter_province').set_value(province).run()
                app.selectbox(key='filter_as_of').set_value(pd.Timestamp(month)).run()
                self.good(app)
                availability, rises = self.kpi_details(app)
                frame = availability.dataframe[0].value
                expected = monthly_findings(changes, province, month)
                self.assertEqual(len(frame), 25)
                self.assertEqual(frame['Nama komoditas'].tolist(), sorted(changes.commodity.unique()))
                self.assertEqual(frame['No.'].tolist(), [str(n) for n in range(1, 26)])
                self.assertEqual((frame['Status data'] == 'Tersedia').sum(), available)
                self.assertTrue((frame.loc[frame['Status data'] == 'Tidak tersedia', 'Harga pada bulan referensi'] == '—').all())
                self.assertEqual(app.metric[0].value, f'{available} / 25')
                self.assertEqual(app.metric[1].value, '—' if count is None else str(count))
                if count:
                    detail = rises.dataframe[0].value
                    self.assertEqual(len(detail), count)
                    self.assertEqual(detail['Nama komoditas'].tolist(), expected['alerts'].commodity.tolist())
                    from analyst.presentation import percent, rupiah
                    self.assertEqual(detail['Perubahan harga (%)'].tolist(), [percent(v, digits=2) for v in expected['alerts'].mom_pct])
                    self.assertEqual(detail['Harga bulan sebelumnya'].tolist(), [rupiah(v) for v in expected['alerts'].previous_price])
                    self.assertEqual(detail['Harga bulan referensi'].tolist(), [rupiah(v) for v in expected['alerts'].price])
                else:
                    self.assertEqual(len(rises.dataframe), 0)
                    message = rises.info[0].value
                    self.assertIn('belum dapat dihitung' if count is None else 'Tidak ada komoditas yang mengalami kenaikan minimal 10%', message)
        app.selectbox(key='filter_province').set_value('Jawa Barat').run()
        before = self.kpi_details(app)[1].dataframe[0].value.copy()
        app.selectbox(key='filter_window').set_value(6).run()
        pd.testing.assert_frame_equal(before, self.kpi_details(app)[1].dataframe[0].value)

    def test_missing_region_does_not_reuse_old_prices(self):
        app=self.open_app()
        app.selectbox(key='filter_province').set_value('Papua Pegunungan').run();self.good(app)
        self.assertEqual(app.metric[0].value,'0 / 25')
        self.assertEqual(app.metric[1].value,'—')

    def test_filters_follow_navigation_without_threshold_control(self):
        app=self.open_app()
        app.selectbox(key='filter_commodity').set_value('Cabai Rawit Merah').run()
        app.selectbox(key='filter_window').set_value(36).run()
        self.menu(app,'Perbandingan provinsi')
        self.assertFalse(any(s.key=='filter_window' for s in app.selectbox))
        self.menu(app,'Ringkasan')
        self.assertEqual(app.selectbox(key='filter_commodity').value,'Cabai Rawit Merah')
        self.assertEqual(app.selectbox(key='filter_window').value,36)
        self.assertFalse(any(s.key=='filter_threshold' for s in app.slider))
        self.assertNotIn('threshold', app.session_state['filters'])
        self.assertEqual(app.metric[1].value,'2')

    def test_selected_trend_changes_and_insufficient_history_is_explained(self):
        app=self.open_app();self.menu(app,'Tren & fluktuasi')
        before=app.get('plotly_chart')[0].proto.spec
        app.selectbox(key='filter_commodity').set_value('Cabai Rawit Merah').run();self.good(app)
        self.assertNotEqual(before,app.get('plotly_chart')[0].proto.spec)
        app.select_slider(key='filter_coverage').set_value(100).run();self.good(app)
        self.assertEqual(app.metric[1].value,'100%')
        app.selectbox(key='filter_as_of').set_value(pd.Timestamp('2021-01-01')).run();self.good(app)
        self.assertEqual(app.metric[0].value,'0')
        self.assertTrue(any('Minimal dua bulan' in i.value for i in app.info))

    def test_regions_toggle_changes_displayed_series(self):
        app=self.open_app();self.menu(app,'Perbandingan provinsi')
        before=app.get('plotly_chart')[0].proto.spec
        app.toggle(key='filter_all_regions').set_value(True).run();self.good(app)
        self.assertNotEqual(before,app.get('plotly_chart')[0].proto.spec)
        columns=[str(c) for df in app.dataframe for c in df.value.columns]
        self.assertNotIn('Satuan asumsi',columns)
        self.assertNotIn('unit',columns)

    def test_report_has_analysis_without_budget_or_assumed_units(self):
        app=self.open_app()
        app.button(key='prepare_excel').click().run();self.good(app)
        book=load_workbook(BytesIO(app.session_state['excel_report'][1]))
        self.assertIn('Bulan pembanding',book.sheetnames)
        self.assertNotIn('Sensitivitas latihan',book.sheetnames)
        self.assertNotIn('Simulasi Anggaran',book.sheetnames)
        self.assertNotIn('Satuan asumsi',[c.value for c in book['Harga Bersih'][1]])
        summary={r[0]:r[1] for r in book['Ringkasan'].iter_rows(min_row=2,values_only=True)}
        self.assertEqual(summary['Judul'],TITLE)
        self.assertFalse(any('keranjang' in str(k).lower() for k in summary))

    def test_old_simulator_session_returns_to_summary(self):
        app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=45)
        app.session_state['page']='Simulator anggaran'
        app.session_state['filters']={'threshold':40}
        app.session_state['filter_threshold']=40
        app.run();self.good(app)
        self.assertEqual(app.radio(key='page').value,'Ringkasan')
        self.assertEqual(app.metric[1].value,'2')
        self.assertNotIn('threshold', app.session_state['filters'])

if __name__=='__main__':unittest.main()
