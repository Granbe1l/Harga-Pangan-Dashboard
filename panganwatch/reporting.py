"""Offline, reproducible Excel and factual executive summaries for PanganWatch."""

from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from panganwatch.analytics import (
    basket_simulation,
    comparison_frame,
    province_comparison,
    volatility_summary,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]

METHOD_NOTES = [
    "Snapshot historis harga konsumen bulanan; bukan harga real-time, transaksi, atau penawaran pemasok.",
    "CSV diperoleh dari salinan pihak ketiga yang mencantumkan Badan Pangan Nasional sebagai sumber asal. Kesesuaian dengan versi terbaru sumber asal belum diverifikasi secara independen.",
    "Satuan kg/liter merupakan asumsi kerja per komoditas; CSV asli tidak menyediakan kolom satuan. Verifikasi sebelum penggunaan operasional.",
    "MoM membandingkan bulan kalender sebelumnya; YoY membandingkan bulan yang sama tahun sebelumnya. Bulan kosong tidak dilewati dan tidak diimputasi.",
    "Volatilitas adalah koefisien variasi harga (simpangan baku sampel/mean × 100), minimal tiga observasi; bukan probabilitas kenaikan atau prediksi.",
    "Median antarwilayah tidak berbobot dan hanya mencakup wilayah dengan data pada bulan/komoditas terpilih; bukan harga nasional resmi atau inflasi IHK.",
    "Seri Papua dan Papua Barat sebelum perubahan skema 2024 dipisahkan dari seri sesudahnya. Jumlah seri geografis tidak sama dengan jumlah provinsi Indonesia pada satu waktu.",
    "Simulasi menggunakan harga tepat pada bulan yang dipilih. Harga tidak tersedia tidak diganti harga bulan lain. Total lengkap tidak diterbitkan bila harga komoditas aktif belum lengkap.",
    "Skenario anggaran berasal dari kuantitas dan perubahan harga yang dimasukkan pengguna; bukan penghematan atau dampak bisnis yang sudah terealisasi.",
    "Lisensi redistribusi snapshot belum terkonfirmasi; periksa ketentuan sumber asal dan penyedia salinan sebelum redistribusi atau penggunaan komersial.",
]


def _window(df: pd.DataFrame, as_of: Any | None, months: int) -> tuple[pd.DataFrame, pd.Timestamp, pd.Timestamp]:
    if isinstance(months, bool) or not isinstance(months, int) or months < 1:
        raise ValueError("Jendela analisis harus berupa jumlah bulan bulat positif.")
    frame = comparison_frame(df)
    if frame.empty:
        raise ValueError("Data harga kosong; laporan tidak dapat dibuat.")
    end = pd.Timestamp(as_of) if as_of is not None else frame["date"].max()
    if pd.isna(end):
        raise ValueError("Bulan laporan tidak valid.")
    end = end.to_period("M").to_timestamp()
    start = end - pd.DateOffset(months=months - 1)
    return frame, start, end


def _records(frame: pd.DataFrame, columns: list[str], n: int = 5) -> list[dict[str, Any]]:
    available = [c for c in columns if c in frame]
    # pandas' JSON writer converts NaN to JSON null and timestamps to ISO dates.
    return json.loads(frame[available].head(n).to_json(orient="records", date_format="iso", double_precision=8))


def build_summary(
    df: pd.DataFrame,
    province: str = "Jawa Barat",
    as_of: Any | None = None,
    months: int = 12,
    threshold_pct: float = 10,
) -> dict[str, Any]:
    """Return JSON-safe, observed findings with exact-month comparisons.

    Rankings are within the selected province. Missing data and absent
    comparisons are explicit, including when the province has no current data.
    """
    if not pd.notna(threshold_pct) or not 0 <= float(threshold_pct) < float("inf"):
        raise ValueError("Ambang kenaikan harus berupa persentase nonnegatif dan terbatas.")
    frame, start, end = _window(df, as_of, months)
    snapshot = frame[(frame["province"] == province) & (frame["date"] == end)]
    selected = frame[frame["province"] == province]
    risk = volatility_summary(selected, start, end)
    increases = snapshot[snapshot["mom_pct"] > 0].sort_values("mom_pct", ascending=False)
    alerts = snapshot[snapshot["mom_pct"] >= float(threshold_pct)].sort_values("mom_pct", ascending=False)
    comparable = snapshot["mom_pct"].notna().sum()
    geography_column = "province_code" if "province_code" in frame else "province"
    return {
        "title": "PanganWatch — Ringkasan Analisis Harga Pangan",
        "province": province,
        "as_of": end.strftime("%Y-%m-%d"),
        "window_start": start.strftime("%Y-%m-%d"),
        "window_months": months,
        "threshold_pct": float(threshold_pct),
        "dataset": {
            "rows": int(len(frame)),
            "geographic_series": int(frame[geography_column].nunique()),
            "commodities": int(frame["commodity"].nunique()),
            "period_start": frame["date"].min().strftime("%Y-%m-%d"),
            "period_end": frame["date"].max().strftime("%Y-%m-%d"),
        },
        "snapshot": {
            "observed_commodities": int(len(snapshot)),
            "mom_comparable_commodities": int(comparable),
            "without_previous_month": int(len(snapshot) - comparable),
            "alert_count": int(len(alerts)),
        },
        "top_increases": _records(increases, ["commodity", "price", "unit", "previous_price", "mom_pct", "yoy_pct"]),
        "alerts": _records(alerts, ["commodity", "price", "unit", "previous_price", "mom_pct"], len(alerts)),
        "top_volatility": _records(risk[risk["cv_pct"].notna()], ["commodity", "cv_pct", "observations", "expected_months", "coverage_pct", "latest_date", "latest_price", "unit"]),
        "methodology": METHOD_NOTES,
    }


DISPLAY = {
    "province_code": "Kode seri wilayah",
    "province": "Provinsi / seri historis",
    "commodity": "Komoditas",
    "date": "Bulan",
    "price": "Harga (Rp)",
    "unit": "Satuan asumsi",
    "source_province_code": "Kode provinsi sumber",
    "geography_note": "Catatan wilayah",
    "unit_assumed": "Satuan diasumsikan",
    "previous_price": "Harga bulan sebelumnya (Rp)",
    "mom_pct": "Perubahan MoM (%)",
    "year_ago_price": "Harga setahun sebelumnya (Rp)",
    "yoy_pct": "Perubahan YoY (%)",
    "observations": "Bulan teramati",
    "expected_months": "Bulan dalam jendela",
    "coverage_pct": "Cakupan data (%)",
    "mean_price": "Rata-rata harga (Rp)",
    "std_price": "Simpangan baku harga (Rp)",
    "cv_pct": "Koefisien variasi (%)",
    "min_price": "Harga minimum (Rp)",
    "max_price": "Harga maksimum (Rp)",
    "latest_date": "Bulan observasi terakhir",
    "latest_price": "Harga terakhir tersedia (Rp)",
    "median_price": "Median wilayah teramati (Rp)",
    "deviation_pct": "Selisih dari median (%)",
    "quantity": "Kuantitas kebutuhan",
    "shock_pct": "Skenario perubahan harga (%)",
    "scenario_price": "Harga skenario (Rp)",
    "baseline_cost": "Anggaran acuan (Rp)",
    "scenario_cost": "Anggaran skenario (Rp)",
    "cost_change": "Perubahan anggaran (Rp)",
    "price_available": "Harga tersedia bulan ini",
}


def _display(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in result.select_dtypes(include="datetime").columns:
        result[column] = result[column].dt.strftime("%Y-%m-%d")
    return result.rename(columns=DISPLAY)


def _metadata() -> pd.DataFrame:
    rows = [{"Keterangan": f"Metode {n}", "Isi": note} for n, note in enumerate(METHOD_NOTES, 1)]
    for filename, label in [("data/source.json", "Provenansi snapshot"), ("data/processed/quality_report.json", "Kualitas snapshot bundel")]:
        path = PROJECT_ROOT / filename
        if not path.exists():
            continue
        content = json.loads(path.read_text(encoding="utf-8"))
        for key, value in content.items():
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            rows.append({"Keterangan": f"{label} — {key}", "Isi": value})
    return pd.DataFrame(rows)


def _summary_table(summary: dict[str, Any]) -> pd.DataFrame:
    data = summary["dataset"]
    snap = summary["snapshot"]
    items = [
        ("Judul", summary["title"]),
        ("Provinsi analisis", summary["province"]),
        ("Bulan laporan", summary["as_of"]),
        ("Awal jendela volatilitas", summary["window_start"]),
        ("Jendela analisis (bulan)", summary["window_months"]),
        ("Ambang peringatan kenaikan (%)", summary["threshold_pct"]),
        ("Observasi pada data masukan", data["rows"]),
        ("Seri geografis pada data masukan (termasuk historis)", data["geographic_series"]),
        ("Komoditas pada data masukan", data["commodities"]),
        ("Awal periode data", data["period_start"]),
        ("Akhir periode data", data["period_end"]),
        ("Komoditas tersedia pada bulan/provinsi terpilih", snap["observed_commodities"]),
        ("Komoditas dengan pembanding MoM", snap["mom_comparable_commodities"]),
        ("Komoditas tanpa pembanding MoM", snap["without_previous_month"]),
        ("Komoditas melewati ambang peringatan", snap["alert_count"]),
    ]
    for index, item in enumerate(summary["top_increases"], 1):
        items.append((f"Kenaikan MoM urutan {index}", f"{item['commodity']}: {item['mom_pct']:.2f}%"))
    for index, item in enumerate(summary["top_volatility"], 1):
        items.append((f"CV harga urutan {index}", f"{item['commodity']}: {item['cv_pct']:.2f}% ({item['observations']}/{item['expected_months']} bulan)"))
    return pd.DataFrame(items, columns=["Metrik", "Nilai"])


def _style_workbook(writer: pd.ExcelWriter) -> None:
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=11, color="20372F")
    header_fill = PatternFill("solid", fgColor="145846")
    header_alignment = Alignment(vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=False)
    wrapped_alignment = Alignment(vertical="top", wrap_text=True)
    for name, sheet in writer.sheets.items():
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30
        headers = [str(cell.value or "") for cell in sheet[1]]
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # Explicit text typing prevents formula injection without
                # changing the source string by prepending an apostrophe.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = body_font
                cell.alignment = wrapped_alignment if name in {"Sumber & Metode", "Ringkasan", "Ringkasan Skenario"} else body_alignment
                label = headers[cell.column - 1]
                if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                    if "(Rp)" in label:
                        cell.number_format = '"Rp" #,##0.00;[Red]-"Rp" #,##0.00'
                    elif "(%)" in label or "Kuantitas" in label:
                        # Values are already percent points: 10 means 10%,
                        # hence do not use Excel's multiplying percent format.
                        cell.number_format = "0.00"
                    else:
                        cell.number_format = "#,##0.##"
        for column_number, label in enumerate(headers, 1):
            if name == "Sumber & Metode":
                width = 49 if column_number == 1 else 110
            elif name in {"Ringkasan", "Ringkasan Skenario"}:
                width = 60 if column_number == 1 else 85
            else:
                values = [len(str(row[column_number - 1].value or "")) for row in sheet.iter_rows(min_row=2, max_row=min(sheet.max_row, 150))]
                width = min(44, max(15, len(label) + 2, max(values, default=0) + 2))
            sheet.column_dimensions[get_column_letter(column_number)].width = width
        if name == "Sumber & Metode":
            for index in range(2, sheet.max_row + 1):
                sheet.row_dimensions[index].height = 45


def workbook_bytes(
    df: pd.DataFrame,
    province: str = "Jawa Barat",
    as_of: Any | None = None,
    months: int = 12,
    basket: pd.DataFrame | None = None,
    threshold_pct: float = 10,
    commodity: str = "Beras Medium",
) -> bytes:
    """Create a styled multi-sheet .xlsx file using local data only."""
    frame, start, end = _window(df, as_of, months)
    summary = build_summary(df, province, end, months, threshold_pct)
    snapshot = frame[(frame["province"] == province) & (frame["date"] == end)].sort_values("commodity")
    risk = volatility_summary(frame[frame["province"] == province], start, end)
    comparison = province_comparison(df, commodity, end)
    comparison["Referensi"] = "Median tidak berbobot; bukan rata-rata nasional resmi"
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _summary_table(summary).to_excel(writer, sheet_name="Ringkasan", index=False)
        _display(df).to_excel(writer, sheet_name="Harga Bersih", index=False)
        _display(snapshot).to_excel(writer, sheet_name="Snapshot Provinsi", index=False)
        _display(risk).to_excel(writer, sheet_name="Risiko Harga", index=False)
        _display(comparison).to_excel(writer, sheet_name="Banding Provinsi", index=False)
        _metadata().to_excel(writer, sheet_name="Sumber & Metode", index=False)
        if basket is not None:
            scenario = basket_simulation(df, province, end, basket)
            _display(scenario["lines"]).to_excel(writer, sheet_name="Simulasi Anggaran", index=False)
            pd.DataFrame([
                ("Status", "Lengkap" if scenario["complete"] else "Belum lengkap / tidak ada kuantitas aktif"),
                ("Harga tidak tersedia", ", ".join(scenario["missing_commodities"]) or "Tidak ada"),
                ("Total anggaran acuan (Rp)", scenario["baseline_total"]),
                ("Total anggaran skenario (Rp)", scenario["scenario_total"]),
                ("Perubahan anggaran (Rp)", scenario["difference"]),
                ("Perubahan anggaran (%)", scenario["difference_pct"]),
                ("Subtotal acuan komoditas tersedia (Rp)", scenario["known_baseline_total"]),
                ("Subtotal skenario komoditas tersedia (Rp)", scenario["known_scenario_total"]),
                ("Catatan", "Skenario pengguna; bukan penghematan aktual. Satuan kg/liter adalah asumsi."),
            ], columns=["Metrik", "Nilai"]).to_excel(writer, sheet_name="Ringkasan Skenario", index=False)
        _style_workbook(writer)
    return output.getvalue()


def summary_markdown(summary: dict[str, Any]) -> str:
    """A factual, shareable Indonesian report; no invented impact claims."""
    data, snap = summary["dataset"], summary["snapshot"]
    lines = [
        "# PanganWatch — Ringkasan Analisis",
        "",
        f"Provinsi: **{summary['province']}**. Bulan laporan: **{summary['as_of'][:7]}**.",
        f"Jendela volatilitas: {summary['window_start'][:7]}–{summary['as_of'][:7]} ({summary['window_months']} bulan).",
        "",
        f"Data masukan: **{data['rows']:,} observasi**, **{data['commodities']} komoditas**, **{data['geographic_series']} seri geografis** termasuk seri historis; periode {data['period_start'][:7]}–{data['period_end'][:7]}.",
        f"Pada bulan/provinsi terpilih tersedia **{snap['observed_commodities']} komoditas**; **{snap['mom_comparable_commodities']}** memiliki pembanding bulan kalender sebelumnya.",
        f"Ada **{snap['alert_count']} komoditas** dengan kenaikan MoM ≥ {summary['threshold_pct']:g}%. Ambang ini adalah parameter analisis, bukan batas resmi.",
        "",
        "## Kenaikan harga bulanan yang teramati",
        "",
    ]
    if summary["top_increases"]:
        for item in summary["top_increases"]:
            lines.append(f"- **{item['commodity']}**: naik **{item['mom_pct']:.2f}%**, dari Rp{item['previous_price']:,.0f} menjadi Rp{item['price']:,.0f} per {item.get('unit', 'satuan asumsi')}.")
    else:
        lines.append("Tidak ada kenaikan positif dengan pembanding MoM yang tersedia pada pilihan ini.")
    lines += ["", "## Variasi harga dalam jendela analisis", ""]
    if summary["top_volatility"]:
        for item in summary["top_volatility"]:
            lines.append(f"- **{item['commodity']}**: CV **{item['cv_pct']:.2f}%**, {item['observations']}/{item['expected_months']} bulan teramati (cakupan {item['coverage_pct']:.1f}%).")
    else:
        lines.append("Belum ada seri dengan minimal tiga observasi untuk menghitung CV.")
    lines += [
        "",
        "## Implikasi dan batas penggunaan",
        "",
        "Komoditas dengan kenaikan atau variasi tinggi dapat diprioritaskan untuk evaluasi anggaran. Uji skenario kuantitas dan harga di simulator sebelum menyusun rekomendasi pembelian. Data ini belum cukup untuk membuktikan penghematan, menentukan pemasok termurah, atau memprediksi kenaikan berikutnya.",
        "",
    ]
    lines += [f"- {note}" for note in summary["methodology"]]
    return "\n".join(lines) + "\n"
