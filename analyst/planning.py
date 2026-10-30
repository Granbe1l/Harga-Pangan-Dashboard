"""Budget decisions with explicit completeness and price provenance."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import math

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from analyst.analytics import basket_simulation, _month

DEFAULT_BASKET = {
    "Beras Medium": 25.0, "Daging Ayam Ras": 10.0, "Telur Ayam Ras": 5.0,
    "Minyak Goreng Curah": 5.0, "Cabai Rawit Merah": 2.0, "Bawang Merah": 2.0,
}


def default_plan() -> dict:
    return {name: {"quantity": quantity, "shock_pct": 10.0} for name, quantity in DEFAULT_BASKET.items()}


def historical_budget(changes: pd.DataFrame, province: str, month, quantities=None) -> dict:
    """Compare fixed quantities on two exact calendar months, never partial totals."""
    quantities = DEFAULT_BASKET if quantities is None else quantities
    if any(not math.isfinite(float(q)) or q < 0 for q in quantities.values()):
        raise ValueError("Kuantitas harus nonnegatif dan terbatas.")
    requested = pd.DataFrame([{"commodity": c, "quantity": q} for c, q in quantities.items() if q > 0], columns=["commodity", "quantity"])
    month = _month(month)
    current = changes[(changes.province == province) & (changes.date == month)]
    previous = changes[(changes.province == province) & (changes.date == month - pd.DateOffset(months=1))]
    lines = requested.merge(current[["commodity", "price", "unit"]], on="commodity", how="left", validate="one_to_one")
    # Look up both months independently: missing current rows must not erase valid prior prices.
    lines = lines.merge(previous[["commodity", "price"]].rename(columns={"price":"previous_price"}), on="commodity", how="left", validate="one_to_one")
    units = changes.drop_duplicates("commodity").set_index("commodity").unit
    lines["unit"] = lines["unit"].fillna(lines.commodity.map(units))
    lines["current_cost"] = lines.quantity * lines.price
    lines["previous_cost"] = lines.quantity * lines.previous_price
    lines["impact"] = lines.current_cost - lines.previous_cost
    available = int(lines.price.notna().sum())
    comparable = int(lines.impact.notna().sum())
    count = len(lines)
    current_total = float(lines.current_cost.sum()) if count and available == count else None
    previous_total = float(lines.previous_cost.sum()) if count and lines.previous_price.notna().all() else None
    delta = current_total - previous_total if current_total is not None and previous_total is not None else None
    return {"lines": lines, "total": current_total, "previous_total": previous_total,
            "delta": delta, "delta_pct": delta / previous_total * 100 if delta is not None and previous_total else None,
            "available": available, "comparable": comparable, "count": count,
            "known_total": float(lines.current_cost.sum()),
            "comparable_delta": float(lines.impact.sum()) if comparable else None}


def apply_edits(plan: dict, names: list[str], edits: dict) -> dict:
    """Apply Streamlit editor deltas by commodity, preserving hidden rows and inputs."""
    result = deepcopy(plan)
    columns = {"Jumlah": "quantity", "Perubahan (%)": "shock_pct", "Harga acuan (Rp)": "override", "Satuan": "unit_override"}
    for index, patch in edits.get("edited_rows", {}).items():
        name = names[int(index)]
        for display, value in patch.items():
            if display in columns:
                if display == "Satuan":
                    if value not in {"kg", "liter", "satuan"}:
                        raise ValueError("Pilih satuan kg, liter, atau satuan.")
                elif value is not None and (not math.isfinite(float(value))):
                    raise ValueError("Isian harus berupa angka terbatas.")
                result[name][columns[display]] = value
    return result


def scenario(data: pd.DataFrame, province: str, month, plan: dict, chosen: list[str], manual=False) -> dict:
    """Use current snapshot unless a row was explicitly edited in manual mode."""
    month = pd.Timestamp(month)
    units = data.drop_duplicates("commodity").set_index("commodity").unit.to_dict()
    snapshot = data[(data.province == province) & (data.date == month)].set_index("commodity").price.to_dict()
    prices, provenance, basket = [], [], []
    for name in chosen:
        row = plan[name]
        overridden = manual and "override" in row
        value = row.get("override") if overridden else snapshot.get(name)
        source = "Belum tersedia" if value is None or pd.isna(value) else ("Diubah pengguna" if overridden else "Snapshot historis")
        user_unit = row.get("unit_override") if manual else None
        provenance.append({"commodity": name, "price_source": source,
            "unit_source": "Dinyatakan pengguna" if user_unit else "Asumsi latihan; belum terverifikasi",
            "unit": user_unit or units.get(name, "satuan")})
        if value is not None and pd.notna(value):
            prices.append({"province": province, "commodity": name, "date": month, "price": value, "unit": units.get(name, "satuan")})
        basket.append({"commodity": name, "quantity": row["quantity"], "shock_pct": row["shock_pct"]})
    reference = pd.DataFrame(prices, columns=["province", "commodity", "date", "price", "unit"])
    result = basket_simulation(reference, province, month, pd.DataFrame(basket))
    result["lines"] = result["lines"].drop(columns="unit").merge(pd.DataFrame(provenance), on="commodity", validate="one_to_one")
    active = result["lines"][result["lines"].quantity > 0]
    result["basis"] = "Input pengguna; belum diaudit" if len(active) and (active.price_source == "Diubah pengguna").all() and (active.unit_source == "Dinyatakan pengguna").all() else "Estimasi latihan; harga atau satuan masih memakai asumsi"
    result["budget_limit"] = None
    return result


def complete_months(data, province, names):
    names = list(names)
    if not names:
        return []
    counts = data[(data.province == province) & data.commodity.isin(names)].groupby("date").commodity.nunique()
    return sorted(counts[counts == len(names)].index, reverse=True)


def scenario_table(result, province, month, period):
    frame = result["lines"].rename(columns={"commodity": "Bahan", "quantity": "Jumlah", "unit": "Satuan",
        "price": "Harga acuan (Rp)", "shock_pct": "Perubahan (%)", "scenario_price": "Harga skenario (Rp)",
        "baseline_cost": "Biaya acuan (Rp)", "scenario_cost": "Biaya skenario (Rp)", "cost_change": "Selisih (Rp)",
        "price_available": "Harga tersedia", "price_source": "Sumber harga", "unit_source": "Dasar satuan"})
    for key, value in reversed([("Wilayah", province), ("Bulan referensi", pd.Timestamp(month).strftime("%Y-%m")), ("Periode belanja", period)]):
        frame.insert(0, key, value)
    return frame


def scenario_workbook(result, province, month, period):
    """Export the actual edited scenario, including row-level sources and missing totals."""
    table = scenario_table(result, province, month, period)
    summary = pd.DataFrame({"Keterangan": ["Jenis hasil", "Wilayah", "Bulan referensi", "Periode belanja", "Status", "Anggaran acuan", "Anggaran skenario", "Selisih", "Subtotal acuan tersedia", "Subtotal skenario tersedia", "Dasar perhitungan", "Batas anggaran", "Sisa anggaran"],
        "Nilai": ["Skenario asumsi, bukan pengeluaran terealisasi", province, pd.Timestamp(month).strftime("%Y-%m"), period,
            "Lengkap" if result["complete"] else "Belum lengkap", result["baseline_total"], result["scenario_total"], result["difference"], result["known_baseline_total"], result["known_scenario_total"], result["basis"], result.get("budget_limit"), result.get("headroom")]})
    notes = pd.DataFrame({"Catatan": ["Belanja acuan = jumlah × harga acuan yang digunakan.",
        "Harga skenario = harga acuan × (1 + perubahan / 100).", "Sumber harga dicatat per bahan; harga yang tidak tersedia tetap kosong.",
        "Dasar satuan dicatat per baris: asumsi latihan atau pernyataan pengguna; belum diaudit. Tanpa ongkir, susut, biaya operasional, stok, atau margin.",
        "Kuantitas ditentukan pengguna per periode belanja; bukan catatan transaksi usaha."]})
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in [("Ringkasan", summary), ("Rincian bahan", table), ("Metode", notes)]:
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="102A43")
                cell.font = Font(color="FFFFFF", bold=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        cell.number_format = '#,##0.00'
            for i, col in enumerate(ws.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(65, max(18, max(len(str(c.value or "")) for c in col) + 2))
    return buffer.getvalue()


def basket_coverage(data, province, quantities=None):
    """Count full prices and valid monthly comparisons for one fixed basket."""
    quantities = DEFAULT_BASKET if quantities is None else quantities
    names = [name for name, amount in quantities.items() if amount > 0]
    dates = pd.date_range(data.date.min(), data.date.max(), freq="MS")
    selected = data[(data.province == province) & data.commodity.isin(names)]
    counts = selected.groupby("date").commodity.nunique().reindex(dates, fill_value=0)
    complete = counts.eq(len(names)) & bool(names)
    prior = complete.reindex(dates - pd.DateOffset(months=1), fill_value=False).to_numpy()
    comparisons = complete & prior
    return {"months": len(dates), "complete_months": int(complete.sum()), "comparison_months": int(comparisons.sum()),
        "first_complete": complete[complete].index.min() if complete.any() else None,
        "last_complete": complete[complete].index.max() if complete.any() else None,
        "calendar": pd.DataFrame({"date": dates, "available": counts.to_numpy(), "complete": complete.to_numpy(), "comparable": comparisons.to_numpy()})}


def sensitivity_cases(changes, province, month):
    """Transparent counterfactual, never presented as an observed business basket."""
    cases = []
    for label, egg in [("Latihan A: telur 5 kg",5.0),("Latihan B: telur 1 kg",1.0),("Latihan C: lima bahan tanpa minyak",5.0)]:
        quantities = dict(DEFAULT_BASKET)
        quantities["Telur Ayam Ras"] = egg
        if label.startswith("Latihan C"): quantities.pop("Minyak Goreng Curah")
        result = historical_budget(changes, province, month, quantities)
        observed = result["lines"].dropna(subset=["impact"])
        positive = observed[observed.impact > 0].sort_values("impact", ascending=False)
        cases.append({"Profil latihan":label,"Kontributor terbesar": positive.iloc[0].commodity if len(positive) else "Belum ada kenaikan yang dapat dibandingkan",
            "Dampak bahan (Rp)":float(positive.iloc[0].impact) if len(positive) else None,"Perubahan total (Rp)":result["delta"]})
    return pd.DataFrame(cases)


def apply_budget_limit(result, limit):
    """Never compare a partial subtotal with an entire spending limit."""
    if not math.isfinite(float(limit)) or limit < 0:
        raise ValueError("Batas anggaran harus nonnegatif dan terbatas.")
    result = deepcopy(result)
    result["budget_limit"] = float(limit) if limit > 0 else None
    result["headroom"] = float(limit) - result["scenario_total"] if limit > 0 and result["complete"] else None
    return result
