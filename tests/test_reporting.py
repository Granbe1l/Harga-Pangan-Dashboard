"""Workbook content, injection protection, and factual summary invariants."""
from io import BytesIO
import json
import unittest

from openpyxl import load_workbook
import pandas as pd

from analyst.reporting import build_summary, summary_markdown, workbook_bytes


def fixture():
    rows = []
    for date, price in [("2024-01-01", 10000), ("2024-02-01", 12000), ("2024-03-01", 15000)]:
        rows.append({"province_code": "32", "province": "Jawa Barat", "commodity": "Beras Medium", "date": date, "price": float(price), "unit": "kg"})
    rows.append({"province_code": "33", "province": "Jawa Tengah", "commodity": "Beras Medium", "date": "2024-03-01", "price": 13000.0, "unit": "kg"})
    return pd.DataFrame(rows)


class ReportingTests(unittest.TestCase):
    def test_factual_summary_and_json_serialization(self):
        summary = build_summary(fixture(), as_of="2024-03", months=3, threshold_pct=20)
        json.dumps(summary, allow_nan=False)
        self.assertEqual(summary["dataset"]["rows"], 4)
        self.assertEqual(summary["snapshot"]["alert_count"], 1)
        self.assertAlmostEqual(summary["top_increases"][0]["mom_pct"], 25)
        self.assertEqual(summary["top_volatility"][0]["observations"], 3)
        self.assertIn("25.00%", summary_markdown(summary))

    def test_workbook_contains_data_and_methodology(self):
        workbook = load_workbook(BytesIO(workbook_bytes(fixture(), as_of="2024-03", months=3)))
        self.assertEqual(workbook["Harga Bersih"].max_row, 5)
        self.assertEqual(workbook["Snapshot Provinsi"].freeze_panes, "A2")
        methods = " ".join(str(cell.value) for row in workbook["Sumber & Metode"] for cell in row)
        self.assertIn("belum diverifikasi", methods)
        self.assertIn("kg/liter", methods)
        self.assertIn("tidak diimputasi", methods)
        snapshot = workbook["Snapshot Provinsi"]
        headers = {cell.value: cell.column for cell in snapshot[1]}
        mom = snapshot.cell(2, headers["Perubahan MoM (%)"])
        self.assertAlmostEqual(mom.value, 25)
        self.assertNotIn("%", mom.number_format)

    def test_missing_prices_stay_missing_in_scenario_totals(self):
        basket = pd.DataFrame({"commodity": ["Cabai"], "quantity": [2], "shock_pct": [20]})
        workbook = load_workbook(BytesIO(workbook_bytes(fixture(), as_of="2024-03", basket=basket)))
        scenario = workbook["Ringkasan Skenario"]
        values = {row[0].value: row[1].value for row in scenario.iter_rows(min_row=2)}
        self.assertEqual(values["Harga tidak tersedia"], "Cabai")
        self.assertIsNone(values["Total anggaran acuan (Rp)"])
        self.assertIsNone(values["Total anggaran skenario (Rp)"])

    def test_valid_scenario_totals_are_readable(self):
        basket = pd.DataFrame({"commodity": ["Beras Medium"], "quantity": [2], "shock_pct": [10]})
        workbook = load_workbook(BytesIO(workbook_bytes(fixture(), as_of="2024-03", basket=basket)))
        values = {row[0].value: row[1].value for row in workbook["Ringkasan Skenario"].iter_rows(min_row=2)}
        self.assertEqual(values["Total anggaran acuan (Rp)"], 30000)
        self.assertEqual(values["Total anggaran skenario (Rp)"], 33000)

    def test_strings_cannot_become_excel_formulas(self):
        data = fixture()
        data.loc[0, "commodity"] = '=HYPERLINK("https://example.com","x")'
        workbook = load_workbook(BytesIO(workbook_bytes(data, as_of="2024-03")))
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
        self.assertEqual(workbook["Harga Bersih"]["C2"].value, '=HYPERLINK("https://example.com","x")')

    def test_no_comparison_is_unknown_in_report_not_zero_alerts(self):
        summary = build_summary(fixture(), as_of="2024-01")
        self.assertIsNone(summary["snapshot"]["alert_count"])
        self.assertIn("belum dapat dihitung", summary_markdown(summary))
        workbook = load_workbook(BytesIO(workbook_bytes(fixture(), as_of="2024-01")))
        values = {r[0].value:r[1].value for r in workbook["Ringkasan"].iter_rows(min_row=2)}
        self.assertIsNone(values["Komoditas melewati ambang peringatan"])

    def test_missing_calendar_month_does_not_generate_findings(self):
        data = fixture().query("date != '2024-02-01'")
        summary = build_summary(data, as_of="2024-03")
        self.assertEqual(summary["snapshot"]["mom_comparable_commodities"], 0)
        self.assertEqual(summary["top_increases"], [])
        self.assertEqual(summary["top_volatility"], [])


if __name__ == "__main__":
    unittest.main()
