"""Business invariants: historical vs hypothetical, completeness, provenance, exports."""
from io import BytesIO
import unittest
import pandas as pd
from openpyxl import load_workbook
from analyst.analytics import comparison_frame
from analyst.planning import historical_budget, scenario, apply_edits, scenario_workbook, complete_months, basket_coverage, apply_budget_limit, sensitivity_cases


def fixture():
    return pd.DataFrame([
        {"province":"Jawa Barat","commodity":"Beras","date":pd.Timestamp("2025-09-01"),"price":100.,"unit":"kg"},
        {"province":"Jawa Barat","commodity":"Beras","date":pd.Timestamp("2025-10-01"),"price":120.,"unit":"kg"},
        {"province":"Jawa Barat","commodity":"Telur","date":pd.Timestamp("2025-10-01"),"price":200.,"unit":"kg"},
    ])


class PlanningTests(unittest.TestCase):
    def test_historical_impact_uses_fixed_quantity_and_exact_month(self):
        r=historical_budget(comparison_frame(fixture()),"Jawa Barat","2025-10",{"Beras":2})
        self.assertEqual(r['total'],240)
        self.assertEqual(r['delta'],40)
        self.assertAlmostEqual(r['delta_pct'],20)

    def test_incomplete_comparison_never_becomes_total_change(self):
        r=historical_budget(comparison_frame(fixture()),"Jawa Barat","2025-10",{"Beras":2,"Telur":1})
        self.assertEqual(r['total'],440)
        self.assertIsNone(r['delta'])
        self.assertEqual(r['comparable_delta'],40)
        self.assertEqual(r['comparable'],1)

    def test_missing_current_price_has_only_labelled_subtotal(self):
        r=historical_budget(comparison_frame(fixture()),"Jawa Barat","2025-09",{"Beras":2,"Telur":1})
        self.assertIsNone(r['total'])
        self.assertEqual(r['known_total'],200)
        self.assertIsNone(r['delta'])

    def test_zero_price_change_is_real_zero_not_missing(self):
        frame=fixture(); frame.loc[1,'price']=100
        r=historical_budget(comparison_frame(frame),"Jawa Barat","2025-10",{"Beras":2})
        self.assertEqual(r['delta'],0)
        self.assertEqual(r['delta_pct'],0)

    def test_manual_price_provenance_is_per_row_and_does_not_modify_data(self):
        data=fixture(); before=data.copy(deep=True)
        plan={'Beras':{'quantity':2,'shock_pct':10,'override':150},'Telur':{'quantity':1,'shock_pct':0}}
        r=scenario(data,'Jawa Barat','2025-10',plan,list(plan),True)
        sources=r['lines'].set_index('commodity').price_source.to_dict()
        self.assertEqual(sources,{'Beras':'Diubah pengguna','Telur':'Snapshot historis'})
        self.assertEqual(r['baseline_total'],500)
        pd.testing.assert_frame_equal(data,before)
        self.assertEqual(scenario(data,'Jawa Barat','2025-10',plan,list(plan),False)['baseline_total'],440)

    def test_manual_blank_does_not_fall_back_to_snapshot(self):
        p={'Beras':{'quantity':2,'shock_pct':10,'override':None}}
        r=scenario(fixture(),'Jawa Barat','2025-10',p,list(p),True)
        self.assertIsNone(r['baseline_total'])
        self.assertEqual(r['lines'].price_source.iloc[0],'Belum tersedia')

    def test_edits_preserve_other_rows_and_fractional_quantities(self):
        plan={'Beras':{'quantity':2,'shock_pct':10},'Telur':{'quantity':1,'shock_pct':0}}
        edited=apply_edits(plan,['Beras'],{'edited_rows':{0:{'Jumlah':2.5,'Harga acuan (Rp)':130}}})
        self.assertEqual(edited['Telur'],plan['Telur'])
        self.assertEqual(edited['Beras']['quantity'],2.5)
        self.assertEqual(edited['Beras']['override'],130)
        self.assertNotIn('override',plan['Beras'])

    def test_complete_months_require_all_requested_materials(self):
        self.assertEqual(complete_months(fixture(),'Jawa Barat',['Beras','Telur']),[pd.Timestamp('2025-10-01')])
        self.assertEqual(complete_months(fixture(),'Jawa Barat',['Tidak ada']),[])

    def test_scenario_excel_exports_sources_and_blanks_not_false_totals(self):
        p={'Beras':{'quantity':2,'shock_pct':10,'override':None}}
        r=scenario(fixture(),'Jawa Barat','2025-10',p,list(p),True)
        book=load_workbook(BytesIO(scenario_workbook(r,'Jawa Barat','2025-10','Per minggu')))
        summary={row[0]:row[1] for row in book['Ringkasan'].iter_rows(min_row=2,values_only=True)}
        self.assertIsNone(summary['Anggaran acuan'])
        self.assertEqual(summary['Periode belanja'],'Per minggu')
        self.assertIn('Sumber harga',[c.value for c in book['Rincian bahan'][1]])

    def test_previous_total_survives_missing_current_row(self):
        data=fixture();r=historical_budget(comparison_frame(data),'Jawa Barat','2025-11',{'Beras':2})
        self.assertIsNone(r['total'])
        self.assertEqual(r['previous_total'],240)
        self.assertIsNone(r['delta'])

    def test_coverage_counts_gaps_without_inventing_prices(self):
        data=fixture();r=basket_coverage(data,'Jawa Barat',{'Beras':2,'Telur':1})
        self.assertEqual(r['complete_months'],1)
        self.assertEqual(r['comparison_months'],0)
        self.assertEqual(r['months'],2)

    def test_budget_limit_rejects_partial_comparison(self):
        p={'Beras':{'quantity':2,'shock_pct':10,'override':None}}
        result=apply_budget_limit(scenario(fixture(),'Jawa Barat','2025-10',p,list(p),True),1000)
        self.assertIsNone(result['headroom'])
        p['Beras']['override']=100
        result=apply_budget_limit(scenario(fixture(),'Jawa Barat','2025-10',p,list(p),True),250)
        self.assertAlmostEqual(result['headroom'],30)

    def test_price_override_without_unit_declaration_remains_training(self):
        p={'Beras':{'quantity':2,'shock_pct':10,'override':150}}
        result=scenario(fixture(),'Jawa Barat','2025-10',p,list(p),True)
        self.assertIn('latihan',result['basis'])
        p['Beras']['unit_override']='kg'
        result=scenario(fixture(),'Jawa Barat','2025-10',p,list(p),True)
        self.assertEqual(result['lines'].unit_source.iloc[0],'Dinyatakan pengguna')
        self.assertEqual(result['basis'],'Input pengguna; belum diaudit')

    def test_export_escapes_formula_in_user_controlled_metadata(self):
        p={'Beras':{'quantity':2,'shock_pct':10}}
        r=scenario(fixture(),'Jawa Barat','2025-10',p,list(p))
        book=load_workbook(BytesIO(scenario_workbook(r,'=SUM(1,2)','2025-10','Per minggu')))
        self.assertFalse(any(c.data_type=='f' for w in book for row in w for c in row))


if __name__=='__main__': unittest.main()
