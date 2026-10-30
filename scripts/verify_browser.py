#!/usr/bin/env python3
"""Optional real-browser acceptance checks and screenshots (requires Playwright/Chromium)."""
from __future__ import annotations
import argparse
from io import BytesIO
import json
from pathlib import Path
import re
import time

from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook
import pandas as pd

TITLE = 'Analisis Harga Pangan dan Simulasi Anggaran Bahan Baku UMKM Kuliner'
MENUS = [('Ringkasan','ringkasan'),('Simulator anggaran','simulator'),('Tren & fluktuasi','risiko'),('Perbandingan wilayah','wilayah')]


def main():
    args=argparse.ArgumentParser(description=__doc__)
    args.add_argument('--url',default='http://127.0.0.1:8502')
    args.add_argument('--chromium',default='/usr/bin/chromium')
    args.add_argument('--output-dir',type=Path,default=Path('/tmp/analyst-browser'))
    opt=args.parse_args();out=opt.output_dir;out.mkdir(parents=True,exist_ok=True)
    report={'target':opt.url,'scope':'local browser acceptance; not proof of public hosting readiness','checks':[],'screenshots':[]}
    errors=[]
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path=opt.chromium,headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
        def ready(page):
            page.get_by_role('heading',name=TITLE,exact=True).wait_for(timeout=45000)
            page.wait_for_timeout(650)
            assert page.get_by_test_id('stException').count()==0
        def settle(page):
            page.wait_for_timeout(650)
            assert page.get_by_test_id('stException').count()==0
        def menu(page,name,mobile=False):
            if mobile:
                page.get_by_test_id('stExpandSidebarButton').click()
            page.get_by_test_id('stSidebar').get_by_text(name,exact=True).click()
            settle(page)
            if mobile:
                sidebar=page.get_by_test_id('stSidebar')
                button=sidebar.get_by_role('button').filter(has_text='keyboard_double_arrow_left')
                button.click()
                settle(page)
        def select(page,label,value):
            box=page.get_by_role('combobox',name=re.compile(re.escape(label)))
            if label=='Pilih komposisi latihan':
                page.get_by_role('tab',name='Dampak keranjang latihan',exact=True).click();settle(page)
            if label=='Pilih komposisi latihan' and box.count()==0:
                page.get_by_text('Profil latihan, sensitivitas, dan cakupan',exact=True).click();settle(page)
            box.click();box.fill(value)
            page.get_by_role('option',name=value,exact=True).click();settle(page)
        def screenshot(page,name,width,height):
            page.set_viewport_size({'width':width,'height':height});settle(page)
            main=page.get_by_test_id('stMain')
            main.evaluate('(e)=>e.scrollTop=0')
            h=main.evaluate('(e)=>e.scrollHeight')
            page.set_viewport_size({'width':width,'height':max(height,min(5500,h+60))});settle(page)
            main.evaluate('(e)=>e.scrollTop=0')
            assert page.evaluate('document.body.scrollWidth<=innerWidth'),name
            page.screenshot(path=str(out/(name+'.png')),full_page=True)
            report['screenshots'].append(name+'.png')
            page.set_viewport_size({'width':width,'height':height})
        def edit_cell(page,ratio,value):
            canvas=page.locator('[data-testid="stDataFrame"] canvas').first
            canvas.scroll_into_view_if_needed();box=canvas.bounding_box()
            page.mouse.dblclick(box['x']+box['width']*ratio,box['y']+52)
            field=page.locator('input.gdg-input:visible')
            field.wait_for(timeout=5000)
            field.fill(value);field.press('Enter');settle(page)
        page=browser.new_page(viewport={'width':1440,'height':950},device_scale_factor=1)
        page.on('pageerror',lambda e: errors.append(str(e)))
        start=time.perf_counter();page.goto(opt.url);ready(page)
        page.get_by_text('Perubahan harga bulan ini',exact=True).wait_for()
        report['local_initial_ready_seconds']=round(time.perf_counter()-start,2)
        expect(page.locator('[data-testid="stMetricValue"]:visible').nth(0)).to_have_text('25 / 25')
        cta=page.get_by_role('button',name='Coba simulasi belanja →')
        assert cta.bounding_box()['y']+cta.bounding_box()['height']<950
        report['checks'].append('desktop first screen contains observed price findings without assumed basket units and a simulator CTA')
        page.screenshot(path=str(out/'ringkasan-first-screen.png'))
        for label,name in MENUS:
            if label!='Ringkasan':menu(page,label)
            screenshot(page,name,1440,950)
        menu(page,'Ringkasan')
        page.get_by_role('tab',name='Dampak keranjang latihan',exact=True).click();settle(page)
        screenshot(page,'keranjang-latihan',1440,950)
        page.get_by_text('Profil latihan, sensitivitas, dan cakupan',exact=True).click();settle(page)
        select(page,'Pilih komposisi latihan','Latihan B')
        expect(page.locator('[data-testid="stMetricValue"]:visible').nth(1)).to_have_text('Rp 3.923')
        screenshot(page,'profil-b-sensitivitas',1440,950)
        select(page,'Bulan referensi','Januari 2021')
        select(page,'Pilih komposisi latihan','Latihan C')
        expect(page.locator('[data-testid="stMetricValue"]:visible').nth(2)).to_have_text('5 / 5 bahan')
        screenshot(page,'profil-c-riwayat-panjang',1440,950)
        select(page,'Pilih komposisi latihan','Latihan A')
        select(page,'Bulan referensi','Oktober 2025')
        report['checks'].append('profile sensitivity changes contributor; explicit five-item profile supports early history without imputation')
        menu(page,'Simulator anggaran')
        edit_cell(page,.43,'30.5')
        expect(page.locator('[data-testid="stMetricValue"]:visible').first).to_have_text('Rp 1.191.214',timeout=10000)
        page.get_by_text('Edit harga pembelian sendiri',exact=True).click();settle(page)
        edit_cell(page,.75,'20000')
        expect(page.locator('[data-testid="stMetricValue"]:visible').first).to_have_text('Rp 1.397.760',timeout=10000)
        menu(page,'Ringkasan');menu(page,'Simulator anggaran')
        expect(page.locator('[data-testid="stMetricValue"]:visible').first).to_have_text('Rp 1.397.760')
        report['checks'].append('real mouse/keyboard edits: fractional quantity and manual price persist across navigation')
        screenshot(page,'simulator-harga-sendiri',1440,950)
        with page.expect_download() as download:
            page.get_by_role('button',name='Unduh skenario CSV',exact=True).click()
        csvpath=out/'browser-skenario.csv';download.value.save_as(csvpath)
        export=pd.read_csv(csvpath)
        assert export.loc[export.Bahan=='Beras Medium','Sumber harga'].iloc[0]=='Diubah pengguna'
        assert export.loc[export.Bahan=='Daging Ayam Ras','Sumber harga'].iloc[0]=='Snapshot historis'
        assert export.loc[export.Bahan=='Beras Medium','Jumlah'].iloc[0]==30.5
        with page.expect_download() as download:
            page.get_by_role('button',name='Unduh skenario Excel',exact=True).click()
        path=out/'browser-skenario.xlsx';download.value.save_as(path)
        book=load_workbook(path)
        assert {'Ringkasan','Rincian bahan','Metode'}<=set(book.sheetnames)
        assert book['Rincian bahan'].max_row==7
        assert 'Dasar satuan' in export.columns
        assert export['Dasar satuan'].str.contains('Asumsi latihan').all()
        page.get_by_role('spinbutton',name='Batas anggaran (Rp, 0 = tanpa batas)').fill('1200000')
        page.get_by_role('spinbutton',name='Batas anggaran (Rp, 0 = tanpa batas)').press('Enter');settle(page)
        assert 'melewati batas anggaran' in page.get_by_test_id('stMain').inner_text()
        report['checks'].append('browser CSV and Excel downloads preserve edited prices and unverified unit status; spending limit evaluated separately')
        page.get_by_role('button',name='Kembalikan contoh awal',exact=True).click();settle(page)
        expect(page.locator('[data-testid="stMetricValue"]:visible').first).to_have_text('Rp 1.118.460')
        select(page,'Bulan referensi','Januari 2021')
        expect(page.locator('[data-testid="stMetricValue"]:visible').first).to_have_text('—')
        screenshot(page,'simulator-belum-lengkap',1440,950)
        page.get_by_role('button',name=re.compile('Gunakan Oktober 2025')).click();settle(page)
        expect(page.locator('[data-testid="stMetricValue"]:visible').first).to_have_text('Rp 1.118.460')
        report['checks'].append('reset and missing-price recovery button restore complete scenario')
        menu(page,'Tren & fluktuasi')
        assert 'Bulan bersama' in page.get_by_test_id('stMain').inner_text()
        assert 'pasangan bulanan yang sama' in page.get_by_test_id('stMain').inner_text()
        report['checks'].append('fluctuation ranking explicitly shares observation months and monthly comparison pairs')
        select(page,'Bulan referensi','Januari 2021')
        screenshot(page,'risiko-riwayat-terbatas',1440,950)
        menu(page,'Ringkasan');page.get_by_role('tab',name='Temuan harga',exact=True).click();settle(page);screenshot(page,'ringkasan-riwayat-terbatas',1440,950)
        select(page,'Bulan referensi','Oktober 2025');select(page,'Wilayah acuan','Papua Pegunungan')
        screenshot(page,'ringkasan-tanpa-harga',1440,950)
        select(page,'Wilayah acuan','Jawa Barat');menu(page,'Perbandingan wilayah')
        select(page,'Bahan pilihan','Cabai Rawit Merah')
        page.get_by_text('Lihat seluruh wilayah',exact=True).click();settle(page)
        screenshot(page,'wilayah-lengkap',1440,950)
        page.get_by_text('Sumber & cara menghitung',exact=True).click();settle(page)
        screenshot(page,'sumber-metode',1440,950)
        page.get_by_role('button',name='Siapkan laporan Excel',exact=True).click()
        button=page.get_by_role('button',name='Unduh laporan Excel',exact=True);button.wait_for(timeout=60000)
        with page.expect_download() as download:button.click()
        path=out/'browser-laporan.xlsx';download.value.save_as(path)
        assert load_workbook(path)['Harga Bersih'].max_row==32481
        book = load_workbook(path)
        assert 'Sensitivitas latihan' in book.sheetnames
        summary={r[0]:r[1] for r in book['Ringkasan'].iter_rows(min_row=2,values_only=True)}
        assert summary['Bulan keranjang lengkap']==22
        assert summary['Bulan perubahan lengkap']==21
        metadata=' '.join(str(c.value) for row in book['Sumber & Metode'] for c in row)
        assert 'primary_export_user_supplied' in metadata
        assert '561054df5d8b36e7d75d92715fe037031a7e9c7c9f513fd178322cce491845d1' in metadata
        report['checks'].append('analysis Excel includes 32,480 rows, primary-upload receipt hash, coverage, and sensitivity profiles')
        mobile=browser.new_page(viewport={'width':390,'height':844},device_scale_factor=1,is_mobile=True,has_touch=True)
        mobile.on('pageerror',lambda e: errors.append(str(e)))
        mobile.goto(opt.url);ready(mobile)
        mobile.get_by_text('Perubahan harga bulan ini',exact=True).wait_for()
        for label,name in MENUS:
            if label!='Ringkasan':menu(mobile,label,True)
            screenshot(mobile,name+'-mobile',390,844)
        report['checks'].append('all four menus render on 390 px mobile; sidebar works; no body horizontal overflow')
        # A fresh anonymous context must not inherit the first visitor's choices.
        fresh=browser.new_page(viewport={'width':1280,'height':900})
        fresh.goto(opt.url);ready(fresh)
        expect(fresh.locator('[data-testid="stMetricValue"]:visible').first).to_have_text('25 / 25')
        report['checks'].append('fresh anonymous browser context starts at default data without another session inputs')
        assert not errors, errors
        report['javascript_page_errors']=errors
        browser.close()
    report['status']='passed'
    (out/'browser-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
